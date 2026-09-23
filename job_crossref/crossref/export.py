from __future__ import annotations

import csv
from pathlib import Path

from .applications import CompanyApps
from .models import Match
from .sources.links import links_for

CONTACT_COLUMNS = [
    "priority_score", "company", "application_status", "category", "name", "school", "headline",
    "matched_role", "shared_functions", "currently_at_company", "linkedin_url", "school_evidence",
    "source", "draft_note", "note_chars",
]
COMPANY_COLUMNS = [
    "company", "application_status", "latest_applied", "roles", "nyu_stern_contacts", "nyu_contacts",
    "same_team", "recruiters", "top_contact", "nyu_stern_alumni_link", "nyu_alumni_link", "google_xray_link",
]


def contact_rows(matches: list[Match], groups: dict[str, CompanyApps]) -> list[dict]:
    rows = []
    for m in matches:
        c = m.contact
        rows.append({
            "priority_score": m.score,
            "company": m.company,
            "application_status": groups[m.company_key].best_status,
            "category": m.category,
            "name": c.name,
            "school": "NYU Stern" if c.school == "stern" else "NYU",
            "headline": c.headline,
            "matched_role": m.matched_role,
            "shared_functions": ", ".join(m.shared_functions),
            "currently_at_company": {True: "yes", False: "no", None: "unconfirmed"}[c.currently_at_company],
            "linkedin_url": c.linkedin_url,
            "school_evidence": c.school_evidence,
            "source": c.source,
            "draft_note": m.draft_note,
            "note_chars": len(m.draft_note),
        })
    return rows


def company_rows(matches: list[Match], groups: dict[str, CompanyApps]) -> list[dict]:
    by_company: dict[str, list[Match]] = {}
    for m in matches:
        by_company.setdefault(m.company_key, []).append(m)
    ranked = []
    for key, g in groups.items():
        ms = by_company.get(key, [])
        links = links_for(g.name)
        row = {
            "company": g.name,
            "application_status": g.best_status,
            "latest_applied": g.latest_date.isoformat() if g.latest_date else "",
            "roles": " | ".join(dict.fromkeys(g.roles)),
            "nyu_stern_contacts": sum(m.contact.school == "stern" for m in ms),
            "nyu_contacts": sum(m.contact.school == "nyu" for m in ms),
            "same_team": sum(m.category == "Same team" for m in ms),
            "recruiters": sum(m.category == "Recruiter" for m in ms),
            "top_contact": f"{ms[0].contact.name} ({ms[0].category})" if ms else "",
            "nyu_stern_alumni_link": links["nyu_stern_alumni"],
            "nyu_alumni_link": links["nyu_alumni"],
            "google_xray_link": links["google_xray"],
        }
        ranked.append(((-g.priority, -len(ms), g.name), row))
    return [row for _, row in sorted(ranked, key=lambda t: t[0])]


def write_csv(path: Path, rows: list[dict], columns: list[str]) -> None:
    with open(path, "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=columns)
        w.writeheader()
        w.writerows(rows)


def write_xlsx(path: Path, sheets: dict[str, tuple[list[dict], list[str]]]) -> bool:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font
    except ImportError:
        return False
    wb = Workbook()
    wb.remove(wb.active)
    for title, (rows, columns) in sheets.items():
        ws = wb.create_sheet(title)
        ws.append(columns)
        for cell in ws[1]:
            cell.font = Font(bold=True)
        for r in rows:
            ws.append([r.get(c, "") for c in columns])
        ws.freeze_panes = "A2"
        for col in ws.columns:
            width = max(len(str(cell.value or "")) for cell in col[:200])
            ws.column_dimensions[col[0].column_letter].width = min(max(10, width + 2), 60)
    wb.save(path)
    return True
